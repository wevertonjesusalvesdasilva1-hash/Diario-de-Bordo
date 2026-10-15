# Gera docs/folha-modelo/diario-de-bordo-v3.xlsx: a folha "uma linha = uma peça" em Excel (Frente, Verso e Exemplo),
# com as mesmas cores do PDF, pronta para imprimir em A4 paisagem (uma página por aba). Uso: python3 scripts/folha/gerar_excel_v3.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.page import PageMargins
import os

F = lambda **kw: Font(name='Arial', **kw)

NAVY, HEAD, GRP, GREY, BLUE = '1F2A44', 'E9EDF5', 'D6DEEE', 'F3F4F6', '1D4ED8'
thin, thick = Side(style='thin', color='111111'), Side(style='medium', color='111111')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill('solid', fgColor=c)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1)
top_left = Alignment(horizontal='left', vertical='top', wrap_text=True, indent=1)
W = {'A': 5, 'B': 7.5, 'C': 16, 'D': 14}
TR = [('E', 'F', 'G'), ('H', 'I', 'J'), ('K', 'L', 'M'), ('N', 'O', 'P')]   # 4 trechos: chapa, início, fim
for c, i, f in TR: W[c], W[i], W[f] = 7, 6.6, 6.6
W['Q'], W['R'], W['S'] = 7.5, 10, 21
CODES = ['Mecânica', 'Elétrica', 'Preventiva', 'Aguardando peça', 'Falta de peça', 'Ponte rolante', 'Troca de consumíveis', 'Troca de peça', 'Set-up', 'Tyorei', 'Outros']
def rt(*parts, sz=8.5, color='FF111111'):
    bold, plain = InlineFont(rFont='Arial', sz=sz, b=True, color=color), InlineFont(rFont='Arial', sz=sz, color=color)
    return CellRichText([TextBlock(bold, p[1:]) if p.startswith('*') else TextBlock(plain, p) for p in parts])

def box(ws, rng, value=None, font=None, bg=None, align=center, border=True):
    if ':' not in rng: rng = f'{rng}:{rng}'
    a, z = rng.split(':')
    if a != z: ws.merge_cells(rng)
    c = ws[rng.split(':')[0]]
    if value is not None: c.value = value
    if font: c.font = font
    if bg: c.fill = fill(bg)
    c.alignment = align
    if border:
        for row in ws[rng]:
            for x in row: x.border = B
    return c

def setup(ws):
    for k, v in W.items(): ws.column_dimensions[k].width = v
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False

def header(ws, title, sub, signs=('COORDENADOR', 'CHEFE', 'GERENTE')):
    ws.row_dimensions[1].height = 24; ws.row_dimensions[2].height = 18
    box(ws, 'A1:L2', None, bg=NAVY)
    c = ws['A1']; c.value = rt(f'*{title}', f'      {sub}', sz=13, color='FFFFFFFF'); c.font = F(color='FFFFFF', size=14, bold=False); c.alignment = Alignment(vertical='center', indent=1)
    for rng, t in zip(('M1:N1', 'O1:P1', 'Q1:S1'), signs): box(ws, rng, t, F(size=7, bold=True))
    for rng in ('M2:N2', 'O2:P2', 'Q2:S2'): box(ws, rng, '')

def fields(ws, r, ex=None):
    ws.row_dimensions[r].height = 20
    box(ws, f'A{r}:B{r}', 'MÁQUINA:', F(bold=True, size=9))
    m = '○ 05    ○ 06    ● 09    ○ 10    ○ 11    ○ 12    Outra: ______' if ex else '○ 05    ○ 06    ○ 09    ○ 10    ○ 11    ○ 12    Outra: ______'
    box(ws, f'C{r}:H{r}', m, F(bold=True, size=10), align=left)
    box(ws, f'I{r}:J{r}', 'DATA:', F(bold=True, size=9))
    box(ws, f'K{r}:L{r}', ex and '15/10/2026' or '', F(color=BLUE, size=10))
    box(ws, f'M{r}', 'SETOR:', F(bold=True, size=9))
    box(ws, f'N{r}:O{r}', ex and 'Solda RAW' or '', F(color=BLUE, size=10))
    box(ws, f'P{r}', 'FOLHA', F(bold=True, size=9))
    box(ws, f'Q{r}:S{r}', ex and '1 de 1' or '___ de ___', F(color=BLUE if ex else '111111', size=10))

def section(ws, r, title, hint):
    ws.row_dimensions[r].height = 17
    box(ws, f'A{r}:H{r}', title, F(color='FFFFFF', bold=True, size=10), NAVY, left)
    box(ws, f'I{r}:S{r}', hint, F(color='FFFFFF', size=8), NAVY, Alignment(horizontal='right', vertical='center', indent=1))

def prod_table(ws, r, rows, first=1, data=()):
    ws.row_dimensions[r].height = 16; ws.row_dimensions[r + 1].height = 15
    hf = F(bold=True, size=8)
    for col, t in (('A', 'Nº da\npeça'), ('B', 'DATA'), ('C', 'PEÇA'), ('D', 'MODELO'), ('Q', 'PEÇA\nPRONTA ✓'), ('R', 'CONTINUA\nA PEÇA Nº'), ('S', 'OBSERVAÇÃO')):
        box(ws, f'{col}{r}:{col}{r + 1}', t, hf, HEAD)
    for k, (c, i, f) in enumerate(TR):
        box(ws, f'{c}{r}:{f}{r}', f'TRECHO {k + 1}', F(bold=True, size=9), GRP)
        for col, t in ((c, 'CHAPA'), (i, 'INÍCIO'), (f, 'FIM')): box(ws, f'{col}{r + 1}', t, F(bold=True, size=7.5), HEAD)
    for n in range(rows):
        y = r + 2 + n; ws.row_dimensions[y].height = 21
        d = data[n] if n < len(data) else None
        box(ws, f'A{y}', first + n, F(size=8, color='444444'), GREY)
        vals = {'B': d and d['d'], 'C': d and d['p'], 'D': d and d['o'], 'Q': d and ('✓' if d['ok'] else ''), 'R': d and d.get('cont'), 'S': d and d.get('obs', '')}
        for k, (c, i, f) in enumerate(TR):
            t = d['t'][k] if d and k < len(d['t']) else None
            vals[c], vals[i], vals[f] = (t[0], t[1], t[2]) if t else (None, None, None)
        for col in ['B', 'C', 'D'] + [x for tr in TR for x in tr] + ['Q', 'R', 'S']:
            cell = box(ws, f'{col}{y}', vals.get(col) or None, F(color=BLUE, size=10, bold=col == 'Q'), align=left if col == 'S' else center)
            if col in [tr[1] for tr in TR] + [tr[2] for tr in TR]: cell.number_format = 'hh:mm'
        for c, _, _ in TR: ws[f'{c}{y}'].border = Border(left=thick, right=thin, top=thin, bottom=thin)
    for c, _, _ in TR:
        for y in (r, r + 1): ws[f'{c}{y}'].border = Border(left=thick, right=thin, top=thin, bottom=thin)
    return r + 2 + rows

def stop_table(ws, r, rows, first=1, data=()):
    ws.row_dimensions[r].height = 30
    hf = F(bold=True, size=8)
    cols = [('A', 'A', 'Nº'), ('B', 'B', 'Nº DA\nPEÇA\n(ou —)'), ('C', 'C', 'DATA'), ('D', 'D', 'CHAPA'), ('E', 'F', 'INÍCIO'), ('G', 'H', 'FIM'), ('I', 'I', 'TURNO'), ('J', 'K', 'MOTIVO\n(código 1–11)'), ('L', 'S', 'DESCRIÇÃO DA PARADA')]
    for a, z, t in cols: box(ws, f'{a}{r}:{z}{r}', t, hf, HEAD)
    for n in range(rows):
        y = r + 1 + n; ws.row_dimensions[y].height = 19
        d = data[n] if n < len(data) else None
        v = [first + n] + ([d['pc'], d['d'], d['c'], d['i'], d['f'], d['t'], d['mo'], d['ds']] if d else [None] * 8)
        for (a, z, _), val in zip(cols, v):
            cell = box(ws, f'{a}{y}' if a == z else f'{a}{y}:{z}{y}', val, F(size=8, color='444444') if a == 'A' else F(color=BLUE, size=10), GREY if a == 'A' else None, left if a == 'L' else center)
            if a in ('E', 'G'): cell.number_format = 'hh:mm'
    return r + 1 + rows

def text_block(ws, rng, title, body):
    a, z = rng.split(':'); c0, r0 = a[0], int(a[1:]); c1, r1 = z[0], int(z[1:])
    box(ws, f'{c0}{r0}:{c1}{r0}', title, F(color='FFFFFF', bold=True, size=9), NAVY, left)
    c = box(ws, f'{c0}{r0 + 1}:{c1}{r1}', body, F(size=8.5), align=top_left)
    return c

def bottom_front(ws, r):
    for y in range(r + 1, r + 8): ws.row_dimensions[y].height = 15
    text_block(ws, f'A{r}:G{r + 7}', 'COMO PREENCHER', rt(
        '1. ', '*Uma linha = uma peça.', ' Começou peça nova? Linha nova.\n',
        '2. ', '*Trecho', ' = tempo produzindo, ', '*sem parada no meio', '. Parou a máquina? ', '*Feche o trecho', ' (FIM).\n',
        '3. Anote a parada no ', '*quadro 2', ', com o ', '*Nº da peça', ' e o código do motivo.\n',
        '4. Voltou a produzir a mesma peça? Use o ', '*próximo trecho da mesma linha', '.\n',
        '5. Troca de turno no meio da peça: o próximo operador usa o próximo trecho, com a chapa dele.\n',
        '6. Marque ', '*✓ PRONTA', ' quando a peça terminar. Não terminou? Deixe em branco e continue na próxima folha.\n',
        '7. ', '*Mais de 4 trechos na mesma peça?', ' Continue na ', '*linha de baixo', ', com a mesma peça, e escreva em ', '*CONTINUA A PEÇA Nº', ' o número da linha de cima. ✓ só na última linha.'))
    box(ws, f'H{r}:L{r}', 'CÓDIGOS DE PARADA (coluna MOTIVO)', F(color='FFFFFF', bold=True, size=9), NAVY, left)
    half = (len(CODES) + 1) // 2
    for (a, z), part in ((('H', 'J'), list(enumerate(CODES))[:half]), (('K', 'L'), list(enumerate(CODES))[half:])):
        box(ws, f'{a}{r + 1}:{z}{r + 7}', rt(*[x for k, (i, t) in enumerate(part) for x in (f'*{i + 1}', f'  {t}' + ('\n' if k < len(part) - 1 else ''))], sz=9), align=top_left)
    text_block(ws, f'M{r}:S{r + 7}', 'EXEMPLO: VIGA D51, 13:00 ÀS 17:00, PAROU 14:30–15:30', rt(
        '*Quadro 1 (peça 1): ', 'VIGA · D51 · Trecho 1: 13:00–14:30 · Trecho 2: 15:30–17:00 · ✓\n',
        '*Quadro 2: ', 'Nº da peça 1 · 14:30–15:30 · motivo 4 (aguardando peça)\n\n',
        '*Certo: ', '3 h de produção + 1 h parada = 4 h.\n',
        '*Errado: ', 'um trecho 13:00–17:00 com a parada dentro dele (a hora parada conta duas vezes).'))

def bottom_back(ws, r, prod_ranges, stop_ranges):
    for y in range(r + 1, r + 6): ws.row_dimensions[y].height = 21
    box(ws, f'A{r}:G{r}', 'FECHAMENTO DA FOLHA (contramestre)', F(color='FFFFFF', bold=True, size=9), NAVY, left)
    box(ws, f'H{r}:L{r}', 'PEÇAS', F(color='FFFFFF', bold=True, size=9), NAVY, left)
    box(ws, f'M{r}:S{r}', 'LEMBRETE', F(color='FFFFFF', bold=True, size=9), NAVY, left)
    # totals are formulas: they fill in by themselves when the sheet is typed in Excel
    ph = '+'.join(f"SUMPRODUCT(({sh}!{f}{a}:{f}{z}<>\"\")*MOD({sh}!{f}{a}:{f}{z}-{sh}!{i}{a}:{i}{z},1))" for sh, a, z in prod_ranges for _, i, f in TR)
    sh_ = '+'.join(f"SUMPRODUCT(({sh}!G{a}:G{z}<>\"\")*MOD({sh}!G{a}:G{z}-{sh}!E{a}:E{z},1))" for sh, a, z in stop_ranges)
    rows = [('Horas produzindo (soma dos trechos)', f'={ph}'), ('Horas paradas (soma do quadro 2)', f'={sh_}'), ('Total (24 h no dia, ou as horas dos turnos)', f'=E{r + 1}+E{r + 2}')]
    for k, (t, fml) in enumerate(rows):
        y = r + 1 + k
        box(ws, f'A{y}:D{y}', t, F(size=8.5), align=left)
        c = box(ws, f'E{y}:G{y}', fml, F(color=BLUE, bold=True, size=10)); c.number_format = '[h]:mm'
    pr = '+'.join(f"COUNTA({sh}!Q{a}:Q{z})" for sh, a, z in prod_ranges)
    pc = '+'.join(f"SUMPRODUCT(({sh}!C{a}:C{z}<>\"\")*({sh}!Q{a}:Q{z}=\"\"))" for sh, a, z in prod_ranges)
    for k, (t, fml) in enumerate([('Peças prontas (nº de ✓)', f'={pr}'), ('Continuam na próxima folha', f'={pc}'), ('Conferido por (chapa)', None)]):
        y = r + 1 + k
        box(ws, f'H{y}:J{y}', t, F(size=8.5), align=left); box(ws, f'K{y}:L{y}', fml, F(color=BLUE, bold=True, size=10))
    box(ws, f'M{r + 1}:S{r + 3}', rt('*Trecho é só produção.', ' Parou? Feche o trecho e anote a parada com o ', '*Nº da peça', '. Voltou? ', '*Próximo trecho da mesma linha', '. ✓ quando a peça ficar pronta.'), F(size=8.5), align=top_left)

EX_P = [
    {'d': '15/10', 'p': 'LANÇA', 'o': 'PC210-10M0', 't': [['1001', '06:30', '08:10']], 'ok': 1},
    {'d': '15/10', 'p': 'LANÇA', 'o': 'PC210-10M0', 't': [['1001', '08:10', '09:45'], ['1001', '10:20', '11:05']], 'ok': 1},
    {'d': '15/10', 'p': 'VIGA PC', 'o': 'D51-24', 't': [['1001', '11:50', '13:40']], 'ok': 1},
    {'d': '15/10', 'p': 'VIGA PC', 'o': 'D51-24', 't': [['1001', '13:40', '15:00'], ['1002', '15:00', '15:25']], 'ok': 1, 'obs': 'Troca de turno no meio da peça'},
    {'d': '15/10', 'p': 'LANÇA', 'o': 'PC360LC-8M2', 't': [['1002', '16:40', '18:30'], ['1002', '18:50', '20:40']], 'ok': 1},
    {'d': '15/10', 'p': 'LANÇA', 'o': 'PC360LC-8M2', 't': [['1002', '21:10', '23:30']], 'ok': 1},
    {'d': '16/10', 'p': 'BOOM WA', 'o': 'WA200-6', 't': [['1003', '00:20', '03:10']], 'ok': 1},
    {'d': '16/10', 'p': 'BOOM WA', 'o': 'WA200-6', 't': [['1003', '03:40', '06:00']], 'ok': 0, 'obs': 'Continua na próxima folha'},
]
EX_P[4] = {'d': '15/10', 'p': 'LANÇA', 'o': 'PC360LC-8M2', 't': [['1002', '16:40', '17:20'], ['1002', '17:30', '18:00'], ['1002', '18:10', '18:30'], ['1002', '18:50', '19:30']], 'ok': 0, 'obs': '4 trechos cheios: segue na linha 6'}
EX_P[5] = {'d': '15/10', 'p': 'LANÇA', 'o': 'PC360LC-8M2', 't': [['1002', '19:40', '20:40']], 'ok': 1, 'cont': 5, 'obs': 'Continuação da peça 5'}
EX_P.insert(6, {'d': '15/10', 'p': 'LANÇA', 'o': 'PC360LC-8M2', 't': [['1002', '21:10', '23:30']], 'ok': 1})
EX_S = [
    {'pc': 1, 'd': '15/10', 'c': '1001', 'i': '06:00', 'f': '06:30', 't': 1, 'mo': 9, 'ds': 'Set-up do dispositivo da lança'},
    {'pc': 2, 'd': '15/10', 'c': '1001', 'i': '09:45', 'f': '10:20', 't': 1, 'mo': 7, 'ds': 'Troca de bico e arame'},
    {'pc': 3, 'd': '15/10', 'c': '1001', 'i': '11:05', 'f': '11:50', 't': 1, 'mo': 6, 'ds': 'Ponte ocupada no setor'},
    {'pc': '—', 'd': '15/10', 'c': '1002', 'i': '15:25', 'f': '16:40', 't': 2, 'mo': 5, 'ds': 'Sem peça no setor'},
    {'pc': 5, 'd': '15/10', 'c': '1002', 'i': '17:20', 'f': '17:30', 't': 2, 'mo': 7, 'ds': 'Troca de bico'},
    {'pc': 5, 'd': '15/10', 'c': '1002', 'i': '18:00', 'f': '18:10', 't': 2, 'mo': 6, 'ds': 'Ponte para virar a peça'},
    {'pc': 5, 'd': '15/10', 'c': '1002', 'i': '18:30', 'f': '18:50', 't': 2, 'mo': 8, 'ds': ''},
    {'pc': 5, 'd': '15/10', 'c': '1002', 'i': '19:30', 'f': '19:40', 't': 2, 'mo': 10, 'ds': 'Tyorei'},
    {'pc': 7, 'd': '15/10', 'c': '1002', 'i': '20:40', 'f': '21:10', 't': 2, 'mo': 1, 'ds': 'Tocha travando, mecânico ajustou'},
    {'pc': 8, 'd': '15/10', 'c': '1003', 'i': '23:30', 'f': '00:20', 't': 3, 'mo': 9, 'ds': 'Set-up para o boom'},
    {'pc': 9, 'd': '16/10', 'c': '1003', 'i': '03:10', 'f': '03:40', 't': 3, 'mo': 2, 'ds': 'Queda de energia no painel'},
]
def as_time(rows, keys):   # 'hh:mm' → Excel time, so the totals add up
    out = []
    for r in rows:
        r = dict(r)
        if 't' in r and isinstance(r['t'], list): r['t'] = [[c, tm(i), tm(f)] for c, i, f in r['t']]
        for k in keys: r[k] = tm(r[k])
        out.append(r)
    return out
def tm(s):
    from datetime import time
    h, m = map(int, s.split(':')); return time(h % 24, m)

FP, FS, BP, BS = 10, 11, 13, 11
wb = Workbook()
def front(ws, ex=False):
    setup(ws); header(ws, 'DIÁRIO DE BORDO · UMA LINHA = UMA PEÇA', 'EXEMPLO PREENCHIDO · dados fictícios' if ex else 'Trecho = tempo SÓ de produção, sem parada no meio')
    fields(ws, 3, ex)
    section(ws, 4, '1. PRODUÇÃO — cada linha é UMA peça', 'parou no meio da peça? feche o trecho, anote a parada no quadro 2 com o Nº da peça e continue no próximo trecho')
    r = prod_table(ws, 5, FP, 1, as_time(EX_P, []) if ex else ())
    section(ws, r, '2. PARADAS DA MÁQUINA', 'escreva o Nº da peça que estava na máquina (ou “—” se não havia peça)')
    r = stop_table(ws, r + 1, FS, 1, as_time(EX_S, ['i', 'f']) if ex else ())
    bottom_front(ws, r + 1)
    ws.print_area = f'A1:S{r + 8}'
    return 7, 7 + FP - 1, r - FS, r - 1   # production rows, stop rows

def back(ws, name_front, fr, ex=False):
    setup(ws); header(ws, 'DIÁRIO DE BORDO · CONTINUAÇÃO (VERSO)', 'Continue a numeração da frente', ('MÁQUINA', 'DATA', 'FOLHA'))
    if ex:
        for rng, v in zip(('M2:N2', 'O2:P2', 'Q2:S2'), ('RAW09', '15/10/2026', '1 de 1')): ws[rng.split(':')[0]].value = v; ws[rng.split(':')[0]].font = F(color=BLUE, size=10)
    section(ws, 3, '1. PRODUÇÃO (continuação)', 'uma linha = uma peça')
    r = prod_table(ws, 4, BP, FP + 1)
    section(ws, r, '2. PARADAS (continuação)', 'Nº da peça que estava na máquina, ou “—”')
    r2 = stop_table(ws, r + 1, BS, FS + 1)
    t = ws.title.replace("'", "")
    bottom_back(ws, r2 + 1, [(f"'{name_front}'", fr[0], fr[1]), (f"'{t}'", 6, 6 + BP - 1)], [(f"'{name_front}'", fr[2], fr[3]), (f"'{t}'", r + 2, r2 - 1)])
    ws.print_area = f'A1:S{r2 + 4}'

f = wb.active; f.title = 'Frente'; fr = front(f)
back(wb.create_sheet('Verso'), 'Frente', fr)
e = wb.create_sheet('Exemplo (frente)'); er = front(e, True)
back(wb.create_sheet('Exemplo (verso)'), 'Exemplo (frente)', er, True)
out = os.path.join(os.path.dirname(__file__), '..', '..', 'docs', 'folha-modelo', 'diario-de-bordo-v3.xlsx')
wb.save(out); print('ok', os.path.abspath(out))
