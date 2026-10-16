"""Gera o Excel de CONTROLE (PRODUÇÃO + PARADAS) para a folha "uma linha = uma peça".

Uso:  python3 scripts/folha/gerar_controle.py raw  <controle-antigo.xlsx> <saida.xlsx>
      python3 scripts/folha/gerar_controle.py usi  <controle-antigo.xlsx> <saida.xlsx>
      ... acrescentar `historico` no fim para copiar também as linhas de PRODUÇÃO e PARADAS do controle antigo (cada linha antiga vira
          uma linha nova com até 3 trechos; Nº da peça fica em branco; o motivo da parada vai como nome, sem código).

O controle antigo só serve para copiar a aba MATRIZ TEMPO (tempos padrão). Nada dele fica no repositório.
As fórmulas seguem as do controle antigo (COD, tempo de cada trecho, TEMPO/DEC, TEMPO PADRÃO BAAN, TTL, TEMPO/DEC da parada),
agora com 8 trechos por peça e o motivo da parada vindo do código 1–10 da folha.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

kind, src_path, out_path = sys.argv[1:4]
WITH_HIST = len(sys.argv) > 4 and sys.argv[4] == 'historico'
USI = kind == 'usi'
MACH = 'MAQUINA' if USI else 'ROBÔ'
MACHINES = ['UMC 01', 'UMC 02', 'HMC 16', 'NCBZ'] if USI else ['RAW05', 'RAW06', 'RAW09', 'RAW10', 'RAW11', 'RAW12']
CODES = [(1, 'MECÂNICA', 'Máquina ou ferramenta quebrada, trocador, vazamento'), (2, 'ELÉTRICA', 'Falha elétrica, sensor, alarme'),
         (3, 'PREVENTIVA', 'Manutenção preventiva'), (4, 'TROCA DE PEÇA', 'Tirou a pronta e pôs outra igual'),
         (5, 'FALTA DE PEÇA', 'Sem peça para produzir'), (6, 'REFEIÇÃO', 'Almoço, janta, ceia, café'),
         (7, 'TYOREI', 'Tyorei'), (8, 'OUTROS', 'O que não está na lista (descreva)')]
if USI:
    CODES += [(9, 'SET-UP', 'Mudou de modelo (dispositivo, programa)'), (10, 'CONSUMIVEIS', 'Inserto, fresa, broca')]
NT = 8                                   # trechos por peça
ROWS_P, ROWS_S = 1500, 3000              # linhas já preparadas com fórmula (além do histórico, quando houver)

ARIAL = 'Arial'
YEL, BLUE, RED = PatternFill('solid', fgColor='FFFF00'), PatternFill('solid', fgColor='00B0F0'), PatternFill('solid', fgColor='FF0000')
GREY = PatternFill('solid', fgColor='D9D9D9')
thin = Side(style='thin', color='999999')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
F9, F9B = Font(name=ARIAL, size=9), Font(name=ARIAL, size=9, bold=True)
FW = Font(name=ARIAL, size=9, bold=True, color='FFFFFF')


def head(ws, row, col, text, fill=None, font=F9B, h=None):
    c = ws.cell(row, col, text)
    c.font, c.alignment, c.border = font, CEN, BOX
    if fill: c.fill = fill
    return c


HP, HS = [], []                          # histórico copiado do controle antigo
if WITH_HIST:
    old = openpyxl.load_workbook(src_path)
    ow = old['PRODUÇÃO']                       # controle antigo: B máquina, C data, D peça, E modelo, F-H / I-K / L-N os 3 blocos, O turno, P observações
    for r in range(3, ow.max_row + 1):
        g = lambda c: ow.cell(r, c).value
        if not g(2) or not g(3): continue
        HP.append({'m': g(2), 'd': g(3), 'p': g(4), 'o': g(5), 'tr': [(g(6), g(7), g(8)), (g(9), g(10), g(11)), (g(12), g(13), g(14))], 't': g(15), 'ob': g(16)})
    ow2 = old['PARADAS']
    for r in range(2, ow2.max_row + 1):
        mach = ow2.cell(r, 2).value
        if not mach or str(mach).strip().upper() == 'TOTAL' or not ow2.cell(r, 3).value: continue
        HS.append([ow2.cell(r, c).value for c in range(2, 10)])           # robô, data, chapa, início, fim, turno, motivo, descrição
    ROWS_P, ROWS_S = len(HP) + ROWS_P, len(HS) + ROWS_S
wb = openpyxl.Workbook()

# ------------------------------------------------------------------ PRODUÇÃO
ws = wb.active
ws.title = 'PRODUÇÃO'
C = {}                                   # nome -> número da coluna
names = ['COD', MACH, 'DATA', 'Nº PEÇA', 'PEÇA', 'MODELO']
for n in names: C[n] = len(C) + 1
for k in range(1, NT + 1):
    for p in ('CHAPA', 'INÍCIO', 'FIM'): C[f'{p} {k}'] = len(C) + 1
for n in ('CONTINUA A PEÇA Nº', 'TURNO QUE FINALIZOU', 'OBSERVAÇÕES'): C[n] = len(C) + 1
for k in range(1, NT + 1): C[f'TEMPO T{k}'] = len(C) + 1
for n in ('TEMPO/DEC', 'TEMPO PADRÃO BAAN', 'PARADO NA PEÇA (h)', 'AVISO', 'CHAVE'): C[n] = len(C) + 1
col = lambda n: L(C[n])
last = len(C)
HR, R0, R1 = 2, 3, 2 + ROWS_P            # linha do cabeçalho, primeira e última linha de dados

calc = {'COD', 'TEMPO/DEC', 'TEMPO PADRÃO BAAN', 'PARADO NA PEÇA (h)', 'AVISO', 'CHAVE'} | {f'TEMPO T{k}' for k in range(1, NT + 1)}
for n, c in C.items():
    fill = RED if n in calc else BLUE if n.startswith('CHAPA') else None
    head(ws, HR, c, n, fill, FW if n in calc else F9B)
for k in range(1, NT + 1):
    c0 = C[f'CHAPA {k}']
    ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 2)
    head(ws, 1, c0, f'TRECHO {k}' + (' (1º turno)' if k == 1 else ''), YEL, Font(name=ARIAL, size=11, bold=True))
    for j in (1, 2): ws.cell(1, c0 + j).border = BOX
c0 = C['TEMPO T1']
ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + NT - 1)
head(ws, 1, c0, 'TEMPO DE CADA TRECHO (fim − início)', RED, Font(name=ARIAL, size=11, bold=True, color='FFFFFF'))
ws.merge_cells(start_row=1, start_column=C['TEMPO/DEC'], end_row=1, end_column=C['AVISO'])
head(ws, 1, C['TEMPO/DEC'], 'TOTAIS E CONFERÊNCIA', RED, Font(name=ARIAL, size=11, bold=True, color='FFFFFF'))
ws.row_dimensions[1].height, ws.row_dimensions[2].height = 24, 36

ini = [col(f'INÍCIO {k}') for k in range(1, NT + 1)]
fim = [col(f'FIM {k}') for k in range(1, NT + 1)]
cha = [col(f'CHAPA {k}') for k in range(1, NT + 1)]
tmp = [col(f'TEMPO T{k}') for k in range(1, NT + 1)]
for r in range(R0, R1 + 1):
    f = {}
    f['COD'] = f'=IF({col(MACH)}{r}="","",CONCATENATE(TRIM({col(MACH)}{r})&TRIM({col("PEÇA")}{r})&TRIM({col("MODELO")}{r})))'
    for k in range(NT):
        f[f'TEMPO T{k + 1}'] = f'=IF(AND({ini[k]}{r}<>"",{fim[k]}{r}<>""),MOD({fim[k]}{r}-{ini[k]}{r},1),"")'
    f['TEMPO/DEC'] = f'=IF(COUNT({tmp[0]}{r}:{tmp[-1]}{r})=0,"",SUM({tmp[0]}{r}:{tmp[-1]}{r})*24)'
    f['TEMPO PADRÃO BAAN'] = f"=IF({col('COD')}{r}=\"\",\"\",IFERROR(VLOOKUP({col('COD')}{r},'MATRIZ TEMPO'!A:H,8,0),0))"
    f['PARADO NA PEÇA (h)'] = (f'=IF(OR({col(MACH)}{r}="",{col("Nº PEÇA")}{r}="",{col("DATA")}{r}=""),"",'
                               f'SUMIFS(PARADAS!$L:$L,PARADAS!$A:$A,{col(MACH)}{r},PARADAS!$B:$B,{col("Nº PEÇA")}{r},PARADAS!$D:$D,{col("DATA")}{r}))')
    a = ','.join(f'{c}{r}' for c in ini); b = ','.join(f'{c}{r}' for c in fim)
    f['AVISO'] = (f'=IF({col("COD")}{r}="","",IF(COUNT({a})<>COUNT({b}),"trecho sem início ou sem fim",'
                  f'IF(COUNT({a})=0,"sem horário",IF({col("TEMPO PADRÃO BAAN")}{r}=0,"sem tempo padrão",""))))')
    f['CHAVE'] = f'=IF(OR({col(MACH)}{r}="",{col("Nº PEÇA")}{r}="",{col("DATA")}{r}=""),"",{col(MACH)}{r}&"|"&{col("DATA")}{r}&"|"&{col("Nº PEÇA")}{r})'
    h = HP[r - R0] if r - R0 < len(HP) else None
    if h:
        f[MACH], f['DATA'], f['PEÇA'], f['MODELO'] = h['m'], h['d'], h['p'], h['o']
        for k, (c_, i_, f_) in enumerate(h['tr'], start=1):
            f[f'CHAPA {k}'], f[f'INÍCIO {k}'], f[f'FIM {k}'] = c_, i_, f_
        f['TURNO QUE FINALIZOU'], f['OBSERVAÇÕES'] = h['t'], h['ob']
    for n, c in C.items():
        cell = ws.cell(r, c)
        if n in f: cell.value = f[n]
        cell.font = F9
        cell.alignment = Alignment(horizontal='left' if n in ('OBSERVAÇÕES', 'AVISO') else 'center', vertical='center')
        cell.number_format = ('dd/mm/yyyy' if n == 'DATA' else 'h:mm' if n.startswith(('INÍCIO', 'FIM', 'TEMPO T')) else '0.00' if n in ('TEMPO/DEC', 'TEMPO PADRÃO BAAN', 'PARADO NA PEÇA (h)') else 'General')
widths = {'COD': 16, MACH: 9, 'DATA': 11, 'Nº PEÇA': 7, 'PEÇA': 24, 'MODELO': 16, 'CONTINUA A PEÇA Nº': 11, 'TURNO QUE FINALIZOU': 11, 'OBSERVAÇÕES': 34,
          'TEMPO/DEC': 10, 'TEMPO PADRÃO BAAN': 11, 'PARADO NA PEÇA (h)': 11, 'AVISO': 26, 'CHAVE': 12}
for n, c in C.items():
    ws.column_dimensions[L(c)].width = widths.get(n, 9 if n.startswith('CHAPA') else 7.5)
ws.column_dimensions[col('CHAVE')].hidden = True
for k in range(1, NT + 1): ws.column_dimensions[L(C[f'TEMPO T{k}'])].outlineLevel = 1
ws.freeze_panes = ws.cell(R0, C['MODELO'] + 1)
t = Table(displayName='Tabela1', ref=f'A{HR}:{L(last)}{R1}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
ws.add_table(t)
dv = DataValidation(type='list', formula1='"' + ','.join(MACHINES) + '"', allow_blank=True, showErrorMessage=True, errorTitle='Máquina', error='Escolha uma máquina da lista.')
dv.add(f'{col(MACH)}{R0}:{col(MACH)}{R1}'); ws.add_data_validation(dv)
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='3', allow_blank=True, showErrorMessage=True, errorTitle='Turno', error='Turno 1, 2 ou 3.')
dv.add(f'{col("TURNO QUE FINALIZOU")}{R0}:{col("TURNO QUE FINALIZOU")}{R1}'); ws.add_data_validation(dv)
av = f'{col("AVISO")}{R0}:{col("AVISO")}{R1}'
ws.conditional_formatting.add(av, FormulaRule(formula=[f'{col("AVISO")}{R0}<>""'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))

# ------------------------------------------------------------------ PARADAS
sp = wb.create_sheet('PARADAS')
S = ['ROBÔ' if not USI else 'MAQUINA', 'Nº DA PEÇA', 'TURNO', 'DATA', 'CHAPA', 'INICIO', 'FINAL', 'CÓD', 'MOTIVO DA PARADA', 'DESCRIÇÃO DA PARADA', 'TTL', 'TEMPO/DEC', 'PEÇA']
SC = {n: i + 1 for i, n in enumerate(S)}
scalc = {'MOTIVO DA PARADA', 'TTL', 'TEMPO/DEC', 'PEÇA'}
for n, c in SC.items():
    head(sp, 1, c, n, YEL if n not in scalc else RED, F9B if n not in scalc else FW)
sp.row_dimensions[1].height = 30
M, NP, TU, DA, CH, I, F, CO, MO, DE, TT, TD, PE = [L(i) for i in range(1, 14)]
SR1 = 1 + ROWS_S
# a aba PRODUÇÃO consulta estas colunas de PARADAS: A máquina, B Nº da peça, D data, L TEMPO/DEC
assert (M, NP, DA, TD) == ('A', 'B', 'D', 'L')
for r in range(2, SR1 + 1):
    f = {
        'MOTIVO DA PARADA': f'=IF({CO}{r}="","",IFERROR(VLOOKUP({CO}{r},CÓDIGOS!$A:$B,2,0),"CÓDIGO?"))',
        'TTL': f'=IF(OR({I}{r}="",{F}{r}=""),"",MOD({F}{r}-{I}{r},1))',
        'TEMPO/DEC': f'=IF({TT}{r}="","",{TT}{r}*24)',
        'PEÇA': (f'=IF({NP}{r}="","",IF(OR({NP}{r}="—",{NP}{r}="-"),"(sem peça)",IFERROR(INDEX(PRODUÇÃO!${col("PEÇA")}:${col("PEÇA")},'
                 f'MATCH({M}{r}&"|"&{DA}{r}&"|"&{NP}{r},PRODUÇÃO!${col("CHAVE")}:${col("CHAVE")},0)),"?")))'),
    }
    if r - 2 < len(HS):
        m_, d_, c_, i_, f_, t_, mo_, ds_ = HS[r - 2]
        f[S[0]], f['DATA'], f['CHAPA'], f['INICIO'], f['FINAL'], f['TURNO'], f['DESCRIÇÃO DA PARADA'] = m_, d_, c_, i_, f_, t_, ds_
        f['MOTIVO DA PARADA'] = str(mo_).strip() if mo_ else None            # histórico: o nome do motivo, sem código
    for n, c in SC.items():
        cell = sp.cell(r, c)
        if n in f: cell.value = f[n]
        cell.font = F9
        cell.alignment = Alignment(horizontal='left' if n in ('DESCRIÇÃO DA PARADA', 'PEÇA', 'MOTIVO DA PARADA') else 'center', vertical='center')
        cell.number_format = 'dd/mm/yyyy' if n == 'DATA' else 'h:mm' if n in ('INICIO', 'FINAL', 'TTL') else '0.00' if n == 'TEMPO/DEC' else 'General'
for n, w in zip(S, [9, 9, 7, 11, 9, 8, 8, 6, 20, 46, 8, 10, 24]): sp.column_dimensions[L(SC[n])].width = w
sp.freeze_panes = 'A2'
t = Table(displayName='Tabela2', ref=f'A1:{L(len(S))}{SR1}')
t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
sp.add_table(t)
dv = DataValidation(type='list', formula1='"' + ','.join(MACHINES) + '"', allow_blank=True, showErrorMessage=True, errorTitle='Máquina', error='Escolha uma máquina da lista.')
dv.add(f'A2:A{SR1}'); sp.add_data_validation(dv)
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='3', allow_blank=True, showErrorMessage=True, errorTitle='Turno', error='Turno 1, 2 ou 3.')
dv.add(f'{TU}2:{TU}{SR1}'); sp.add_data_validation(dv)
dv = DataValidation(type='whole', operator='between', formula1='1', formula2=str(len(CODES)), allow_blank=True, showErrorMessage=True, errorTitle='Código', error=f'Código de 1 a {len(CODES)} (veja a aba CÓDIGOS).')
dv.add(f'{CO}2:{CO}{SR1}'); sp.add_data_validation(dv)
sp.conditional_formatting.add(f'{PE}2:{PE}{SR1}', FormulaRule(formula=[f'{PE}2="?"'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))
sp.conditional_formatting.add(f'{MO}2:{MO}{SR1}', FormulaRule(formula=[f'{MO}2="CÓDIGO?"'], fill=PatternFill('solid', bgColor='FFC7CE'), font=Font(color='9C0006')))

# ------------------------------------------------------------------ CÓDIGOS
cs = wb.create_sheet('CÓDIGOS')
for i, h in enumerate(['CÓD', 'MOTIVO (vai para a aba PARADAS)', 'Quando usar']): head(cs, 1, i + 1, h, YEL)
for i, (c, m, d) in enumerate(CODES, start=2):
    for j, v in enumerate((c, m, d), start=1):
        x = cs.cell(i, j, v); x.font = F9; x.border = BOX
        x.alignment = Alignment(horizontal='center' if j == 1 else 'left')
cs.column_dimensions['A'].width, cs.column_dimensions['B'].width, cs.column_dimensions['C'].width = 7, 32, 52
cs.cell(len(CODES) + 3, 1, 'Pode mudar o nome do motivo na coluna B: a aba PARADAS acompanha.').font = Font(name=ARIAL, size=9, italic=True)

# ------------------------------------------------------------------ MATRIZ TEMPO (copiada do controle antigo)
src = openpyxl.load_workbook(src_path)['MATRIZ TEMPO']
srcv = openpyxl.load_workbook(src_path, data_only=True)['MATRIZ TEMPO']       # valores calculados (para células que apontam para outra aba)
mx = wb.create_sheet('MATRIZ TEMPO')
for i, h in enumerate(['', 'SEQ.', 'NÚMERO ITEM ', 'PEÇA ', 'FOTO', 'MÁQUINA', 'MODELO', 'TEMPO PADRÃO BAAN', 'TEMPO POR MINUTOS BAAN', 'obs']):
    if i: head(mx, 1, i + 1, h, YEL)
head(mx, 1, 1, 'COD', RED, FW)
n = 0
for r in range(2, src.max_row + 1):
    vals = [src.cell(r, c).value for c in range(1, 11)]
    if not vals[3] and not vals[5] and not vals[6]: continue                  # linha vazia
    n += 1
    mx.cell(r, 1, f'=CONCATENATE(TRIM(F{r})&TRIM(D{r})&TRIM(G{r}))')
    for c in (2, 3, 4, 6, 7, 9, 10):
        v = vals[c - 1]
        if isinstance(v, str) and v.startswith('=') and '!' in v: v = srcv.cell(r, c).value     # fórmula que apontava para outra aba: fica o valor
        mx.cell(r, c, v)
    h = vals[7]
    if isinstance(h, str) and h.startswith('='):
        h = f'=IFERROR({h[1:].replace(chr(39) + "MATRIZ TEMPO" + chr(39) + "!", "")},0)'          # tempo escrito como texto ("SEM TEMPO NO BAAN") vira 0, sem #VALUE!
    mx.cell(r, 8, h)
    for c in range(1, 11): mx.cell(r, c).font = F9
    mx.cell(r, 8).number_format = '0.00'
for c, w in zip('ABCDEFGHIJ', [34, 6, 18, 28, 6, 10, 16, 12, 12, 24]): mx.column_dimensions[c].width = w
mx.freeze_panes = 'A2'

# ------------------------------------------------------------------ COMO USAR
gl = wb.create_sheet('COMO USAR')
text = [
    ('COMO PREENCHER', True),
    ('Uma linha da aba PRODUÇÃO = uma peça da folha (quadro 1). Copie do papel: Nº da peça, data, peça, modelo e os trechos (chapa, início, fim).', False),
    ('Até 8 trechos por peça. A peça que parou 7 vezes cabe numa linha só; não precisa linha de baixo.', False),
    ('Escreva só a máquina, a data e o que está no papel. As colunas vermelhas são fórmulas: não digite nelas.', False),
    ('Uma linha da aba PARADAS = uma parada da folha (quadro 2). Digite o CÓD do motivo (1 a %d): o nome do motivo aparece sozinho.' % len(CODES), False),
    ('', False),
    ('O QUE AS FÓRMULAS FAZEM', True),
    ('COD (PRODUÇÃO): máquina + peça + modelo, para achar o tempo padrão na aba MATRIZ TEMPO.', False),
    ('TEMPO T1 a T8: duas colunas de horário viram o tempo do trecho (fim − início; passa da meia-noite sem problema).', False),
    ('TEMPO/DEC: soma dos trechos em horas decimais = tempo total produzindo a peça (a parada fica fora, porque já foi tirada do trecho).', False),
    ('TEMPO PADRÃO BAAN: tempo padrão da peça na matriz (0 = a peça/modelo não está na matriz).', False),
    ('PARADO NA PEÇA (h): soma das paradas da aba PARADAS com a mesma máquina, data e Nº da peça.', False),
    ('AVISO: trecho com início sem fim (ou fim sem início) e peça sem tempo padrão.', False),
    ('PARADAS: TTL = fim − início; TEMPO/DEC = TTL em horas; PEÇA = a peça da PRODUÇÃO com a mesma máquina, data e Nº (aparece ? se não achar).', False),
    ('', False),
    ('ATENÇÃO', True),
    ('Para a parada achar a peça, a máquina, a data e o Nº da peça precisam estar iguais nas duas abas. Parada sem peça na máquina: escreva — no Nº.', False),
    ('A MATRIZ TEMPO veio do controle que você enviou. Se a matriz mudou, cole os valores novos nas colunas B a J.', False),
    ('As fórmulas já estão preparadas em %d linhas de produção e %d de paradas.' % (ROWS_P, ROWS_S), False),
]
for i, (tx, b) in enumerate(text, start=1):
    c = gl.cell(i, 1, tx); c.font = Font(name=ARIAL, size=11 if b else 10, bold=b); c.alignment = Alignment(wrap_text=True, vertical='top')
gl.column_dimensions['A'].width = 130
wb.move_sheet('COMO USAR', offset=-(len(wb.sheetnames) - 1))
wb.active = 1
for w in wb.worksheets:
    w.sheet_view.tabSelected = (w.title == 'PRODUÇÃO')
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
wb.calculation.fullCalcOnLoad = True
wb.save(out_path)
print('ok', out_path, 'matriz', n, 'colunas produção', last, 'histórico:', len(HP), 'produções', len(HS), 'paradas')
